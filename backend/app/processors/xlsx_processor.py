"""XLSX processor. Walks every sheet in the workbook and produces
column-anchored content units the same way the CSV processor does, plus
one whole-sheet table unit per sheet — so "Sheet Financials, Row 18,
Column Revenue" is something the system can actually point to.
"""
import openpyxl

from .base import DocumentProcessor
from ..pipeline.normalized import NormalizedDocument, ContentUnit, Location


class XLSXProcessor(DocumentProcessor):
    def process(self, file_path: str) -> NormalizedDocument:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        units = []

        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(c) if c is not None else "" for c in rows[0]]
            data_rows = rows[1:]

            for row_idx, row in enumerate(data_rows, start=2):
                for col_idx, header in enumerate(headers):
                    if col_idx >= len(row):
                        continue
                    value = row[col_idx]
                    if value is None or str(value).strip() == "":
                        continue
                    units.append(ContentUnit(
                        location=Location(sheet=sheet.title, row=row_idx, column=header or f"Col {col_idx + 1}"),
                        text=f"{header}: {value} (Sheet {sheet.title}, row {row_idx})",
                        content_type="table",
                    ))

            table_rows = [[str(c) if c is not None else "" for c in r] for r in rows]
            units.append(ContentUnit(
                location=Location(sheet=sheet.title, table=sheet.title),
                text="\n".join(", ".join(r) for r in table_rows),
                tables=[{"name": sheet.title, "rows": table_rows}],
                content_type="table",
            ))

        return NormalizedDocument(units=units, unit_count_label="sheets")
