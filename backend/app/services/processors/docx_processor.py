from docx import Document as DocxDocument

from app.services.processors.base import BaseProcessor, ProcessingResult, RawContentUnit


class DOCXProcessor(BaseProcessor):
    def process(self, file_path: str) -> ProcessingResult:
        doc = DocxDocument(file_path)
        units: list[RawContentUnit] = []
        current_section = "Document"
        para_idx = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style and para.style.name and para.style.name.lower().startswith("heading"):
                current_section = text
                continue
            para_idx += 1
            units.append(RawContentUnit(
                content_type="text",
                text=text,
                table_data=None,
                location={"section": current_section, "paragraph": para_idx},
                metadata={"style": para.style.name if para.style else None},
            ))

        for t_idx, table in enumerate(doc.tables):
            rows_data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows_data:
                continue
            headers, rows = rows_data[0], rows_data[1:]
            units.append(RawContentUnit(
                content_type="table",
                text=None,
                table_data={"headers": headers, "rows": rows},
                location={"section": current_section, "table": t_idx + 1},
                metadata={},
            ))

        return ProcessingResult(
            content_units=units,
            used_ocr=False,
            document_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
